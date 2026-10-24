"""Word 检测报告生成(python-docx).

报告结构(与简历中"检测报告模板化一键导出"对应):
    标题
    一、任务概览      —— 任务号/名称/状态/时间/阈值/图片数/缺陷数(表格)
    二、缺陷类别统计  —— 类别 | 中文显示名 | 数量(表格)
    三、逐图结果      —— 每张图: 标注图插图 + "检出目标 N 个" + 框明细表(表格)

设计: build_task_report 是纯服务函数, 只负责"拿数据 → 生成文件 → 返回路径",
不查数据库(数据由路由层查好传进来), 便于单独测试.
"""
import io

import cv2
import numpy as np
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Inches


def _add_meta_table(doc, task, total_detections: int) -> None:
    """往文档里加"任务概览"两列表格(key | value)."""
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"  # 带边框的内置表格样式
    rows = [
        ("任务号", task.task_no),
        ("任务名称", task.task_name),
        ("状态", str(task.status.value)),
        ("创建时间", str(task.created_at)),
        ("完成时间", str(task.finished_at) if task.finished_at else "-"),
        ("置信度阈值", str(task.confidence_threshold)),
        ("图片总数", str(task.total_images)),
        ("缺陷框总数", str(total_detections)),
    ]
    for k, v in rows:
        cells = table.add_row().cells  # 加一行, 拿到这行的单元格列表
        cells[0].text = k
        cells[1].text = v


def build_task_report(task, images: list, results: list, class_display: dict, base_dir: Path) -> Path:
    """生成检测报告 docx, 返回文件路径.

    参数:
        task          : 任务 ORM 对象
        images        : 该任务的图片 ORM 对象列表
        results       : 该任务的检测框 ORM 对象列表(已按 image_id 排好)
        class_display : {模型类别名: 中文显示名} 映射, 来自 defect_classes 表
        base_dir      : backend/ 目录(报告存到 base_dir/results/{task_no}/report/)
    """
    doc = Document()
    doc.add_heading("焊缝缺陷智能检测报告", 0)  # level=0 是文档大标题

    doc.add_heading("一、任务概览", level=1)
    _add_meta_table(doc, task, len(results))

    doc.add_heading("二、缺陷类别统计", level=1)
    # Counter 统计每个类别出现的框数
    counter = Counter(r.class_name for r in results)
    if counter:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells  # 第一行当表头
        hdr[0].text = "类别"
        hdr[1].text = "显示名"
        hdr[2].text = "数量"
        for name, count in counter.most_common():  # 按数量降序
            cells = table.add_row().cells
            cells[0].text = name
            cells[1].text = class_display.get(name, name)  # 没登记的类别显示原名
            cells[2].text = str(count)
    else:
        doc.add_paragraph("本任务未检出任何目标。")

    doc.add_heading("三、逐图结果", level=1)
    # 把检测框按 image_id 分组: {图片id: [框, 框, ...]}
    results_by_image: dict[int, list] = {}
    for r in results:
        results_by_image.setdefault(r.image_id, []).append(r)

    for img in images:
        doc.add_heading(f"{img.filename}  [{img.status.value}]", level=2)

        # 插图优先级: 标注图 > 增强图 > 原图(前面的还没生成时用后面的兜底)
        rel = img.annotated_path or img.enhanced_path or img.original_path
        abs_path = base_dir / rel
        if abs_path.exists():
            # 中文路径安全读图 -> 压缩成 jpg -> 转 BytesIO 喂给 python-docx
            # (add_picture 不接受裸 bytes, 必须是带 seek() 的文件流 —— 踩过的坑)
            bgr = cv2.imdecode(np.fromfile(str(abs_path), dtype=np.uint8), cv2.IMREAD_COLOR)
            ok, buf = cv2.imencode(".jpg", bgr)
            if ok:
                doc.add_picture(io.BytesIO(buf.tobytes()), width=Inches(5.5))

        boxes = results_by_image.get(img.id, [])
        doc.add_paragraph(f"检出目标 {len(boxes)} 个。")
        if boxes:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "类别"
            hdr[1].text = "置信度"
            hdr[2].text = "框位置 (x1,y1)"
            hdr[3].text = "框位置 (x2,y2)"
            for b in boxes:
                cells = table.add_row().cells
                cells[0].text = class_display.get(b.class_name, b.class_name)
                cells[1].text = f"{float(b.confidence):.4f}"
                cells[2].text = f"({b.x1},{b.y1})"
                cells[3].text = f"({b.x2},{b.y2})"

    # 保存: results/{task_no}/report/{task_no}.docx
    out_dir = base_dir / "results" / task.task_no / "report"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{task.task_no}.docx"
    doc.save(str(out_path))
    return out_path
